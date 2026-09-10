from __future__ import annotations

import hashlib
import io
from datetime import datetime, timezone

from openpyxl import Workbook

from macro_pit.archive import RawArtifact
from macro_pit.db import get_connection, insert_observations
from macro_pit.sources.us_rtdsm import RTDSMSeries, RTDSMSource


def test_rtdsm_wide_vintage_workbook_becomes_long_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["DATE", "RGDP65Q4", "RGDP66Q1"])
    sheet.append(["1965Q3", 100.0, 100.2])
    sheet.append(["1965Q4", 101.0, 101.1])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    raw = tmp_path / "rgdp.xlsx"
    raw.write_bytes(content)
    artifact = RawArtifact(
        source="RTDSM",
        url="https://www.philadelphiafed.org/rgdp.xlsx",
        path=raw.as_posix(),
        sha256=hashlib.sha256(content).hexdigest(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        retrieved_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        size=len(content),
    )
    series = RTDSMSeries("US_REAL_GDP", "RGDP", "Real GDP", "bn_usd_sa_ar", "Q")
    source = RTDSMSource(allow_network=False)
    try:
        rows = source.parse_workbook(content, artifact, series)
    finally:
        source.close()
    assert len(rows) == 4
    assert {row["period"] for row in rows} == {"1965-Q3", "1965-Q4"}
    assert {row["pit_grade"] for row in rows} == {"B"}
    assert rows[0]["available_at"] > rows[0]["release_at"]


def test_numeric_series_prefix_does_not_leak_into_vintage_year(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["DATE", "M165Q4", "M166Q1"])
    sheet.append(["1965M10", 100.0, 100.1])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    raw = tmp_path / "m1.xlsx"
    raw.write_bytes(content)
    artifact = RawArtifact(
        source="RTDSM",
        url="https://www.philadelphiafed.org/m1.xlsx",
        path=raw.as_posix(),
        sha256=hashlib.sha256(content).hexdigest(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        retrieved_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        size=len(content),
    )
    series = RTDSMSeries("US_MONEY_M1", "M1", "M1", "bn_usd_sa", "M")
    source = RTDSMSource(allow_network=False)
    try:
        rows = source.parse_workbook(content, artifact, series)
    finally:
        source.close()
    assert [row["release_at"].year for row in rows] == [1965, 1966]


def test_replaying_full_vintage_matrix_is_idempotent(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["DATE", "RGDP65Q4", "RGDP66Q1"])
    sheet.append(["1965Q3", 100.0, 100.2])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    raw = tmp_path / "rgdp.xlsx"
    raw.write_bytes(content)
    artifact = RawArtifact(
        source="RTDSM", url="https://www.philadelphiafed.org/rgdp.xlsx",
        path=raw.as_posix(), sha256=hashlib.sha256(content).hexdigest(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        retrieved_at=datetime(2026, 9, 1, tzinfo=timezone.utc), size=len(content),
    )
    source = RTDSMSource(allow_network=False)
    try:
        rows = source.parse_workbook(
            content, artifact, RTDSMSeries("US_REAL_GDP", "RGDP", "Real GDP", "bn_usd_sa_ar", "Q")
        )
    finally:
        source.close()
    conn = get_connection(":memory:")
    assert insert_observations(conn, rows).inserted == 2
    replay = insert_observations(conn, rows)
    assert replay.inserted == 0
    assert replay.unchanged == 2
