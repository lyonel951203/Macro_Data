from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import re
from typing import Any, Iterable

from openpyxl import load_workbook

from ..archive import RawArtifact
from ..errors import DataContractError
from .base import BaseSource, SourceInventoryEntry
from .cn_common import extract_period_from_title, extract_release_evidence, html_text, make_observation


class SAFESource(BaseSource):
    source = "SAFE"
    country = "CN"
    parser_version = "safe_release_v2"

    def inventory(self) -> list[SourceInventoryEntry]:
        return [
            SourceInventoryEntry(
                source=self.source,
                dataset="外汇储备、银行结售汇及涉外收付款新闻稿",
                url="https://www.safe.gov.cn/safe/tjsj1/index.html",
                earliest_period=None,
                latest_period=None,
                frequency="M",
                format="html/xls",
                archive_available=True,
                release_timestamp_available=False,
                historical_revision_available=True,
                estimated_count=None,
                evidence="官方统计数据与新闻稿；发布日期精度按原文判级",
            )
        ]

    def parse_release(self, content: bytes, artifact: RawArtifact) -> list[dict]:
        if _is_excel(content, artifact):
            rows = self._parse_bulk_workbook(content, artifact)
            self.validate_row_count(rows)
            return rows

        soup, text = html_text(content)
        if ("官方储备资产" in text or "Official reserve assets" in text) and len(
            re.findall(r"20\d{2}[.]\d{2}", text)
        ) >= 2:
            rows = self._parse_official_reserve_annual(soup, text, artifact)
            self.validate_row_count(rows)
            return rows
        year, month = extract_period_from_title(soup, text)
        release = extract_release_evidence(soup, text, artifact.retrieved_at)
        rows: list[dict] = []

        amount_patterns = [
            ("CN_FX_RESERVE_USD", "FX_RESERVE_USD", "外汇储备", r"外汇储备(?:规模)?(?:为|余额为)?\s*([\d,.]+)\s*亿美元"),
            ("CN_BANK_FX_SETTLEMENT_USD", "BANK_FX_SETTLEMENT_USD", "银行结汇", r"银行结汇\s*([\d,.]+)\s*亿美元"),
            ("CN_BANK_FX_SALES_USD", "BANK_FX_SALES_USD", "银行售汇", r"(?:银行)?售汇\s*([\d,.]+)\s*亿美元"),
            ("CN_CROSS_BORDER_RECEIPTS_USD", "CROSS_BORDER_RECEIPTS_USD", "涉外收入", r"涉外收入\s*([\d,.]+)\s*亿美元"),
            ("CN_CROSS_BORDER_PAYMENTS_USD", "CROSS_BORDER_PAYMENTS_USD", "对外付款", r"对外付款\s*([\d,.]+)\s*亿美元"),
        ]
        values: dict[str, float] = {}
        for canonical, source_id, name, pattern in amount_patterns:
            match = re.search(pattern, text)
            if not match:
                continue
            # One 亿美元 equals 0.1 billion USD.
            value = float(match.group(1).replace(",", "")) / 10.0
            values[canonical] = value
            rows.append(
                make_observation(
                    source=self.source,
                    canonical_series_id=canonical,
                    source_series_id=source_id,
                    series_name=name,
                    unit="bn_usd",
                    year=year,
                    month=month,
                    value=value,
                    artifact=artifact,
                    release=release,
                    parser_version=self.parser_version,
                )
            )
        derived = [
            ("CN_BANK_FX_NET_SETTLEMENT_USD", "BANK_FX_NET_SETTLEMENT_USD", "银行结售汇差额", "CN_BANK_FX_SETTLEMENT_USD", "CN_BANK_FX_SALES_USD"),
            ("CN_CROSS_BORDER_NET_RECEIPTS_USD", "CROSS_BORDER_NET_RECEIPTS_USD", "涉外收付款差额", "CN_CROSS_BORDER_RECEIPTS_USD", "CN_CROSS_BORDER_PAYMENTS_USD"),
        ]
        for canonical, source_id, name, left, right in derived:
            if left not in values or right not in values:
                continue
            rows.append(
                make_observation(
                    source=self.source,
                    canonical_series_id=canonical,
                    source_series_id=source_id,
                    series_name=name,
                    unit="bn_usd",
                    year=year,
                    month=month,
                    value=values[left] - values[right],
                    artifact=artifact,
                    release=release,
                    parser_version=self.parser_version,
                )
            )
        self.validate_row_count(rows)
        return rows

    def _parse_official_reserve_annual(self, soup, text: str, artifact: RawArtifact) -> list[dict]:
        """Parse annual consolidated reserve tables as PIT_C.

        The page date proves when the whole table was available, but not the
        original monthly vintage or first publication time for each cell.
        """
        release = extract_release_evidence(soup, text, artifact.retrieved_at)
        release["pit_grade"] = "C"
        release["release_date_source"] = (
            "official_annual_consolidation:" + str(release["release_date_source"])
        )
        best_periods: list[tuple[int, int]] = []
        best_values: list[float | None] = []
        for table in soup.find_all("table"):
            periods: list[tuple[int, int]] = []
            values: list[float | None] = []
            for tr in table.find_all("tr"):
                cells = [" ".join(cell.get_text(" ", strip=True).split()) for cell in tr.find_all(["th", "td"])]
                found = []
                for cell in cells:
                    period_match = re.fullmatch(r"(20\d{2})[.](0?[1-9]|1[0-2])", cell.strip())
                    if period_match:
                        found.append((int(period_match.group(1)), int(period_match.group(2))))
                if len(found) >= 2:
                    periods = found
                    continue
                if cells and ("外汇储备" in cells[0] or "Foreign currency reserves" in cells[0]):
                    # Each month has an adjacent USD and SDR value. Keep the
                    # first value of every pair, which is USD 100 million.
                    values = [
                        _number(cells[1 + index * 2]) if 1 + index * 2 < len(cells) else None
                        for index in range(len(periods))
                    ]
            if len(periods) > len(best_periods) and values:
                best_periods, best_values = periods, values
        rows: list[dict] = []
        for (year, month), value in zip(best_periods, best_values, strict=True):
            if value is None:
                continue
            rows.append(make_observation(
                source=self.source,
                canonical_series_id="CN_FX_RESERVE_USD",
                source_series_id="FX_RESERVE_USD_ANNUAL_CONSOLIDATION",
                series_name="外汇储备",
                unit="bn_usd",
                year=year,
                month=month,
                value=value / 10.0,
                artifact=artifact,
                release=release,
                parser_version="safe_official_reserve_annual_v1",
            ))
        if len(rows) < 2:
            raise DataContractError("SAFE annual reserve table has fewer than two monthly values")
        return rows

    def _parse_bulk_workbook(self, content: bytes, artifact: RawArtifact) -> list[dict]:
        """Parse official SAFE time-series workbooks as current-history PIT_D.

        These workbooks contain revised historical levels but do not prove what
        was visible on each original release date.  Every row therefore becomes
        available only at the archived file's first-seen timestamp.
        """
        sheets = list(_workbook_sheets(content, artifact))
        release = {
            "release_at": None,
            "available_at": artifact.retrieved_at,
            "release_date_source": "official_bulk_history_first_seen_only",
            "pit_grade": "D",
        }
        rows: list[dict] = []
        for title, grid in sheets:
            searchable = " ".join(
                str(value) for row in grid[:12] for value in row[:12] if value is not None
            ).lower()
            if "foreign exchange reserves" in searchable or "外汇储备" in searchable:
                rows.extend(_reserve_rows(grid, artifact, release, self.parser_version))
            if "settlement and sales" in searchable or "结售汇" in searchable:
                rows.extend(_flow_rows(
                    grid,
                    artifact,
                    release,
                    self.parser_version,
                    settlement_labels=("foreign exchange settlement", "银行结汇"),
                    sales_labels=("foreign exchange sales", "银行售汇"),
                    settlement_id="CN_BANK_FX_SETTLEMENT_USD",
                    sales_id="CN_BANK_FX_SALES_USD",
                    net_id="CN_BANK_FX_NET_SETTLEMENT_USD",
                    names=("银行结汇", "银行售汇", "银行结售汇差额"),
                ))
            if (
                "cross-border receipts and payments" in searchable
                or "international receipts and payments" in searchable
                or "涉外收付款" in searchable
            ):
                rows.extend(_flow_rows(
                    grid,
                    artifact,
                    release,
                    self.parser_version,
                    settlement_labels=("receipts", "涉外收入"),
                    sales_labels=("payments", "对外付款"),
                    settlement_id="CN_CROSS_BORDER_RECEIPTS_USD",
                    sales_id="CN_CROSS_BORDER_PAYMENTS_USD",
                    net_id="CN_CROSS_BORDER_NET_RECEIPTS_USD",
                    names=("涉外收入", "对外付款", "涉外收付款差额"),
                ))
        if not rows:
            raise DataContractError("SAFE workbook has no supported monthly time-series table")
        unique: dict[tuple[str, str], dict] = {}
        for row in rows:
            unique[(row["canonical_series_id"], row["period"])] = row
        return list(unique.values())


def _is_excel(content: bytes, artifact: RawArtifact) -> bool:
    suffix = artifact.path.lower().rsplit(".", 1)[-1]
    return suffix in {"xls", "xlsx"} or content.startswith((b"PK\x03\x04", b"\xd0\xcf\x11\xe0"))


def _workbook_sheets(content: bytes, artifact: RawArtifact) -> Iterable[tuple[str, list[list[Any]]]]:
    if content.startswith(b"PK\x03\x04") or artifact.path.lower().endswith(".xlsx"):
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            max_column = min(sheet.max_column, 500)
            max_row = min(sheet.max_row, 250)
            yield sheet.title, [
                list(row)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_column,
                    values_only=True,
                )
            ]
        return
    try:
        import xlrd
    except ImportError as exc:
        raise DataContractError("reading official SAFE .xls history requires xlrd>=2.0.1") from exc
    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    for sheet in workbook.sheets():
        grid: list[list[Any]] = []
        for row_idx in range(min(sheet.nrows, 250)):
            row: list[Any] = []
            for col_idx in range(min(sheet.ncols, 500)):
                cell = sheet.cell(row_idx, col_idx)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                row.append(value)
            grid.append(row)
        yield sheet.name, grid


def _reserve_rows(
    grid: list[list[Any]],
    artifact: RawArtifact,
    release: dict[str, Any],
    parser_version: str,
) -> list[dict]:
    rows: list[dict] = []
    for row in grid:
        if len(row) < 2:
            continue
        period = _month_value(row[0])
        value = _number(row[1])
        if period is None or value is None:
            continue
        rows.append(make_observation(
            source="SAFE",
            canonical_series_id="CN_FX_RESERVE_USD",
            source_series_id="FX_RESERVE_USD",
            series_name="外汇储备",
            unit="bn_usd",
            year=period.year,
            month=period.month,
            value=value / 10.0,
            artifact=artifact,
            release=release,
            parser_version=parser_version,
        ))
    return rows


def _flow_rows(
    grid: list[list[Any]],
    artifact: RawArtifact,
    release: dict[str, Any],
    parser_version: str,
    *,
    settlement_labels: tuple[str, ...],
    sales_labels: tuple[str, ...],
    settlement_id: str,
    sales_id: str,
    net_id: str,
    names: tuple[str, str, str],
) -> list[dict]:
    header_idx = next(
        (idx for idx, row in enumerate(grid) if sum(_month_value(value) is not None for value in row) >= 2),
        None,
    )
    if header_idx is None:
        return []
    header = grid[header_idx]
    settlement = _find_metric_row(grid[header_idx + 1 :], settlement_labels)
    sales = _find_metric_row(grid[header_idx + 1 :], sales_labels)
    if settlement is None or sales is None:
        return []
    rows: list[dict] = []
    for col_idx, raw_period in enumerate(header):
        period = _month_value(raw_period)
        if period is None:
            continue
        left = _number(settlement[col_idx] if col_idx < len(settlement) else None)
        right = _number(sales[col_idx] if col_idx < len(sales) else None)
        if left is None or right is None:
            continue
        # SAFE monthly workbooks use units of USD 100 million.
        left /= 10.0
        right /= 10.0
        for canonical, source_id, name, value in (
            (settlement_id, settlement_id.removeprefix("CN_"), names[0], left),
            (sales_id, sales_id.removeprefix("CN_"), names[1], right),
            (net_id, net_id.removeprefix("CN_"), names[2], left - right),
        ):
            rows.append(make_observation(
                source="SAFE",
                canonical_series_id=canonical,
                source_series_id=source_id,
                series_name=name,
                unit="bn_usd",
                year=period.year,
                month=period.month,
                value=value,
                artifact=artifact,
                release=release,
                parser_version=parser_version,
            ))
    return rows


def _find_metric_row(grid: list[list[Any]], labels: tuple[str, ...]) -> list[Any] | None:
    normalized_labels = tuple(label.lower() for label in labels)
    for row in grid:
        label = " ".join(str(value) for value in row[:2] if value is not None).strip().lower()
        if any(token in label for token in normalized_labels):
            return row
    return None


def _month_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    text = " ".join(value.strip().split())
    for fmt in ("%B %Y", "%b %Y", "%Y-%m", "%Y/%m", "%Y年%m月"):
        try:
            return datetime.strptime(text, fmt).date().replace(day=1)
        except ValueError:
            continue
    return None


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
            return float(text)
    return None
