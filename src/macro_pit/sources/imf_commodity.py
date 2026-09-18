from __future__ import annotations

import calendar
import io
import math
import re
from datetime import date

from openpyxl import load_workbook

from ..archive import RawArtifact
from .base import BaseSource, SourceInventoryEntry


IMF_SERIES_ID = "PALLFNF"
IMF_XLSX_URL = (
    "https://www.imf.org/-/media/files/research/"
    "commodityprices/monthly/external-data.xlsx"
)


class IMFCommoditySource(BaseSource):
    """IMF all-commodity index from the official current-history workbook."""

    source = "IMF"
    country = "GLB"
    parser_version = "imf_external_data_xlsx_v1"

    def inventory(self) -> list[SourceInventoryEntry]:
        return [
            SourceInventoryEntry(
                source=self.source,
                dataset="IMF Global Price Index of All Commodities",
                url=IMF_XLSX_URL,
                earliest_period=None,
                latest_period=None,
                frequency="M",
                format="xlsx",
                archive_available=True,
                release_timestamp_available=False,
                historical_revision_available=False,
                estimated_count=None,
                evidence=(
                    "Official IMF monthly workbook. Current-history "
                    "values remain PIT_D; the research view applies its approved "
                    "conservative availability lag without changing the grade."
                ),
            )
        ]

    def parse_workbook(
        self, content: bytes, artifact: RawArtifact
    ) -> list[dict]:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        sheet = workbook["External"]
        rows_iter = sheet.iter_rows(values_only=True)
        headers = list(next(rows_iter))
        try:
            period_column = headers.index("Commodity")
            value_column = headers.index(IMF_SERIES_ID)
        except ValueError as exc:
            workbook.close()
            raise ValueError(
                "IMF workbook is missing Commodity/PALLFNF columns"
            ) from exc

        rows: list[dict] = []
        for values in rows_iter:
            raw_period = values[period_column]
            match = re.fullmatch(
                r"(\d{4})M(\d{1,2})", str(raw_period or "")
            )
            if match is None:
                continue
            try:
                value = float(values[value_column])
            except (TypeError, ValueError):
                continue
            if not math.isfinite(value):
                continue
            year, month = int(match.group(1)), int(match.group(2))
            period_end = date(
                year, month, calendar.monthrange(year, month)[1]
            )
            rows.append(
                {
                    "country": self.country,
                    "source": self.source,
                    "canonical_series_id": (
                        "GLB_IMF_ALL_COMMODITY_PRICE_INDEX"
                    ),
                    "source_series_id": IMF_SERIES_ID,
                    "series_name": (
                        "IMF Global Price Index of All Commodities"
                    ),
                    "frequency": "M",
                    "unit": "index_2016_100",
                    "seasonal_adjustment": "NSA",
                    "period": f"{year:04d}-{month:02d}",
                    "period_start": date(year, month, 1),
                    "period_end": period_end,
                    "value": value,
                    "release_at": None,
                    "release_date_source": (
                        "imf_current_history_snapshot_estimated_lag_40d"
                    ),
                    "first_seen_at": artifact.retrieved_at,
                    "available_at": artifact.retrieved_at,
                    "pit_grade": "D",
                    "source_url": IMF_XLSX_URL,
                    "raw_file": artifact.path,
                    "raw_sha256": artifact.sha256,
                    "retrieved_at": artifact.retrieved_at,
                    "parser_version": self.parser_version,
                }
            )
        workbook.close()
        self.validate_row_count(rows)
        return rows

