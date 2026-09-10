"""Offline Wind EDB intake. Never infer original release dates or undo transforms."""
from __future__ import annotations

import base64
import calendar
import hashlib
import json
import math
import struct
import zlib
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .errors import DataContractError
from .timeutils import ensure_aware

PARSER_VERSION = "wind-edb-intake/1"

# Reviewed against the submitted workbook. Unknown codes stay in quarantine.
# Values are (canonical ID, exact visible name, source unit, target unit, scale).
MAPPINGS = {
    "M0001383": ("CN_M1_YOY", "中国:M1:同比", "%", "pct_yoy", 1.0),
    "M0001385": ("CN_M2_YOY", "中国:M2:同比", "%", "pct_yoy", 1.0),
    "M5525763": ("CN_TSF_STOCK_YOY", "中国:社会融资规模存量:同比", "%", "pct_yoy", 1.0),
    "M0000606": ("CN_EXPORT_USD", "中国:出口金额:当月值", "亿美元", "bn_usd", 0.1),
    "M0000608": ("CN_IMPORT_USD", "中国:进口金额:当月值", "亿美元", "bn_usd", 0.1),
    "M0000610": ("CN_TRADE_BALANCE_USD", "中国:贸易差额:当月值", "亿美元", "bn_usd", 0.1),
    "M0000607": ("CN_EXPORT_USD_YOY", "中国:出口金额:当月同比", "%", "pct_yoy", 1.0),
    "M0000609": ("CN_IMPORT_USD_YOY", "中国:进口金额:当月同比", "%", "pct_yoy", 1.0),
}


def decode_edb_comment(comment: str) -> dict:
    """Read Wind's base64 local ZIP/deflate stream, including streams without a directory.

    This is data parsing only: no formulas, external links or commands are executed.
    Bound decompression because comments are supplied input.
    """
    try:
        packed = base64.b64decode(comment, validate=True)
        header = struct.unpack("<IHHHHHIIIHH", packed[:30])
        if packed[:4] != b"PK\x03\x04" or header[3] != 8 or header[2] & 1:
            raise ValueError("not an unencrypted deflate stream")
        start = 30 + header[-2] + header[-1]
        decoder = zlib.decompressobj(-15)
        payload = decoder.decompress(packed[start:], 2_000_000)
        if not decoder.eof:
            raise ValueError("truncated or oversized EDB settings")
        if not payload.startswith(b"EDB\x00"):
            raise ValueError("missing EDB signature")
        settings = json.loads(payload[payload.index(b"{"):].decode("utf-8"))
        if not isinstance(settings.get("metric"), list):
            raise ValueError("missing metric definitions")
        return settings
    except (ValueError, KeyError, struct.error, zlib.error) as exc:
        raise DataContractError(f"Cannot verify Wind export settings: {exc}") from exc


def review_workbook(manifest: dict, *, start_period: str = "2005-01") -> dict:
    raw_path = Path(manifest["path"])
    if hashlib.sha256(raw_path.read_bytes()).hexdigest() != manifest["sha256"]:
        raise DataContractError("Wind raw SHA256 mismatch")
    first_seen = ensure_aware(manifest["first_seen_at"])
    retrieved = ensure_aware(manifest["retrieved_at"])
    workbook = load_workbook(raw_path, data_only=False, keep_links=False)
    cached = load_workbook(raw_path, data_only=True, keep_links=False)
    settings_by_sheet, profiles, cells, observations = {}, [], [], []
    global_keys = set()
    try:
        for sheet in workbook:
            if not sheet["A1"].comment:
                raise DataContractError(f"Missing Wind export settings in {sheet.title}")
            settings = decode_edb_comment(sheet["A1"].comment.text)
            settings_by_sheet[sheet.title] = settings
            metrics = {m["code"]: m for m in settings["metric"]}
            if len(metrics) != len(settings["metric"]):
                raise DataContractError("Duplicate Wind metric codes in settings")
            data = list(cached[sheet.title].values)
            labels = [data[j][0] for j in range(1, 6)]
            if labels != ["指标名称", "频率", "单位", "指标ID", "来源"]:
                raise DataContractError(f"Unexpected Wind headers: {labels}")
            blank_to_zero = (settings["exportConfig"].get("blankTo") == 1
                             and str(settings["exportConfig"].get("blankValue")) == "0")
            for column in range(1, sheet.max_column):
                name, freq, unit, code, origin = [data[j][column] for j in range(1, 6)]
                if code not in metrics or not metrics[code].get("extractTag"):
                    raise DataContractError(f"Missing exported metric definition: {code}")
                metric = metrics[code]
                raw_code = metric.get("rawCode", code)
                original = metrics.get(raw_code, metric)
                commands = [x["command"]["name"] for x in metric.get("changeRecord", [])]
                mapping = MAPPINGS.get(code)
                if mapping and (name != mapping[1] or unit != mapping[2] or freq != "月"
                                or metric["unit"] != unit or metric["name"] != name
                                or metric["freq"] != 4):
                    raise DataContractError(f"Reviewed mapping metadata changed for {code}")
                field_cells = []
                for row_no, row in enumerate(data[6:], 7):
                    when, value = row[0], row[column]
                    if not isinstance(when, (date, datetime)):
                        raise DataContractError(f"Invalid Wind date at {sheet.title}!A{row_no}")
                    when = when.date() if isinstance(when, datetime) else when
                    period = when.strftime("%Y-%m")
                    key = (code, period)
                    if key in global_keys:
                        raise DataContractError(f"Duplicate Wind code/period: {key}")
                    global_keys.add(key)
                    reasons = []
                    if period < start_period:
                        reasons.append("before_requested_start")
                    if when > first_seen.date():
                        reasons.append("future_observation_date")
                    if when.day != calendar.monthrange(when.year, when.month)[1]:
                        reasons.append("not_month_end")
                    if commands:
                        reasons.append("transformed_series")
                    if not mapping:
                        reasons.append("unreviewed_mapping")
                    if sheet.cell(row_no, column + 1).data_type == "f":
                        reasons.append("formula_value")
                    numeric = (isinstance(value, (int, float)) and not isinstance(value, bool)
                               and math.isfinite(value))
                    if not numeric:
                        reasons.append("missing_or_nonnumeric")
                    elif value == 0 and blank_to_zero:
                        reasons.append("ambiguous_zero_after_blank_fill")
                    cell = dict(sheet=sheet.title, cell=sheet.cell(row_no, column + 1).coordinate,
                                wind_code=code, raw_wind_code=raw_code, series_name=name,
                                origin=origin, period=period, source_unit=unit, raw_value=value,
                                status="quarantined" if reasons else "accepted",
                                reasons=";".join(reasons),
                                canonical_series_id=mapping[0] if mapping else None)
                    cells.append(cell)
                    field_cells.append(cell)
                    if not reasons:
                        observations.append(dict(
                            country="CN", source="WIND", canonical_series_id=mapping[0],
                            source_series_id=code, series_name=name, frequency="M", unit=mapping[3],
                            seasonal_adjustment=None, period=period,
                            period_start=when.replace(day=1), period_end=when,
                            value=float(value) * mapping[4], release_at=None, release_date_source=None,
                            first_seen_at=first_seen, available_at=first_seen, pit_grade="D",
                            source_url=manifest["url"], raw_file=manifest["path"],
                            raw_sha256=manifest["sha256"], retrieved_at=retrieved,
                            parser_version=PARSER_VERSION,
                        ))
                nonzero = [c for c in field_cells if isinstance(c["raw_value"], (int, float))
                           and c["raw_value"] != 0 and math.isfinite(c["raw_value"])]
                accepted = [c for c in field_cells if c["status"] == "accepted"]
                profiles.append(dict(
                    wind_code=code, raw_wind_code=raw_code, series_name=name, origin=origin,
                    source_unit=unit, source_frequency=freq, original_name=original["name"],
                    original_unit=original["unit"], original_frequency=original["freq"],
                    operations=";".join(commands), canonical_series_id=mapping[0] if mapping else None,
                    first_nonzero_period=min((c["period"] for c in nonzero), default=None),
                    last_nonzero_period=max((c["period"] for c in nonzero), default=None),
                    numeric_zero_count=sum(c["raw_value"] == 0 for c in field_cells),
                    input_cells=len(field_cells), accepted_cells=len(accepted),
                    accepted_first_period=min((c["period"] for c in accepted), default=None),
                    accepted_last_period=max((c["period"] for c in accepted), default=None),
                    quarantined_cells=len(field_cells)-len(accepted),
                    decision="partial_PIT_D" if accepted else "reexport_original",
                ))
    finally:
        workbook.close()
        cached.close()
    return dict(settings=settings_by_sheet, profiles=profiles, cells=cells, observations=observations)
