"""Build the batch-specific intake handoff from saved validation evidence."""
from pathlib import Path
import json

import pandas as pd

OUT = Path("reports/v2/wind_batch1")
summary = json.loads((OUT / "summary.json").read_text(encoding="utf-8"))
profile = pd.read_csv(OUT / "field_profile.csv").fillna("")
receipt = pd.read_csv(OUT / "p0_receipt.csv").fillna("")
accepted = profile[profile.accepted_cells.gt(0)]
pending = receipt[~receipt.this_batch_received.astype(bool)][[
    "canonical_series_id", "download_search_name", "target_unit"]].copy()
codes = {"CN_NEW_RMB_DEPOSITS_YTD": "M0048261", "CN_TSF_STOCK": "M5525755"}
actions = {
    "CN_M0_YOY": "本批未包含；下载百分数原值。",
    "CN_RMB_LOAN_BAL_YOY": "下载人民币贷款余额同比；本批 s3/M0009969 是余额水平，不能直接充当同比。",
    "CN_RMB_DEPOSIT_BAL_YOY": "下载人民币存款余额同比；本批 s2/M0009940 是余额水平，不能直接充当同比。",
    "CN_NEW_RMB_LOANS_YTD": "下载金融统计口径的新增人民币贷款累计；本批 s6/M5206731 是社融分项当月值，不能替代。",
    "CN_NEW_RMB_DEPOSITS_YTD": "重导 M0048261 原始亿元值；当前 s11 已换成美元。",
    "CN_TSF_STOCK": "重导 M5525755 原始万亿元值；当前 s5 已换成美元并改数量级。",
    "CN_TSF_FLOW_YTD": "优先下载年内累计原值；也可重导 M5206730 原始亿元当月值，待逐年月份齐全并核对口径后生成累计。",
}
pending["verified_original_wind_code"] = pending.canonical_series_id.map(codes).fillna("")
pending["action"] = pending.canonical_series_id.map(actions)
pending.to_csv(OUT / "next_download_7.csv", index=False, encoding="utf-8-sig")

coverage_rows = [dict(指标=r.series_name.replace("中国:", ""), Wind代码=r.wind_code,
                     起始月=r.accepted_first_period, 末月=r.accepted_last_period,
                     入库数据期=int(r.accepted_cells)) for r in accepted.itertuples()]

def markdown_table(rows):
    keys = list(rows[0])
    return "| " + " | ".join(keys) + " |\n| " + " | ".join(["---"]*len(keys)) + " |\n" + "\n".join(
        "| " + " | ".join(str(row[k]) for k in keys) + " |" for row in rows)

sections = [
    ("summary", "Executive Summary", "**首批已部分接入：18 个指标中，8 个指标的 1,958 条记录入库。** 数值覆盖从 2005 年起；海关五项每项 260 个月。其余单元格保留原件及排除原因。\n\n"
     "**这批只补当前历史数值。** 入库来源为 WIND、等级为 PIT_D，首次观察时间为 2026 年 9 月 8 日 14:10:31（上海）；发布时间未知，未添加经验 PIT 日期，严格 PIT 宽表不变。"),
    ("quality", "导出设置解释了三类异常", "**内嵌配置确认了空值填 0、换汇和变频，证据置信度高。** 全表 436 个日期、18 个指标，共 7,848 个数值单元格；其中 2,515 个为 0。由于真实零值和缺失填充值无法区分，不能把这 2,515 个零全部解释为真实数据或全部解释为缺失。此问题影响覆盖判断和模型输入，严重程度高。\n\n"
     "10 个 s2—s11 指标含 EXCHANGE，人民币已换成美元；s5 还改过数量级。s7—s9 含 FREQ_UP，原始频率为季度、输出变为月度。这些值不能直接放进目标人民币/月度原始数据序列，严重程度高。已解出原始 Wind 代码；没有按猜测汇率逆算。\n\n"
     "8 个已映射指标的 changeRecord 为空。仅接入 2005 年起、未出现上述变换且非歧义零的原始数值。共 5,890 个单元格留存未接入，包含早于目标起点的记录；排除原因可以重叠，并非全是坏数据。"),
    ("coverage", "海关连续历史已补入，央行仍有局部缺口", "**海关金额及同比均覆盖 2005 年 1 月至 2026 年 8 月。** 金额的源单位是亿美元，入库乘 0.1 转为十亿美元；260 个月的出口减进口与差额最大残差为 0.01 亿美元，在本次 0.02 亿美元的舍入容差内。这是内部一致性检查，不是对官方原始版本的认证。\n\n"
     "M2 同比覆盖 2005 年 1 月至 2026 年 7 月，共 259 个月。M1 同期接入 258 个月，2020 年 1 月的 0 单独保留待核对。社融存量同比接入 141 期：2005—2014 年每年只有 12 月非零记录，2015 年仅四个季末，2016 年起至 2026 年 7 月连续。下面的起点只表示本次接入的第一期，不代表从该月起连续。"),
    ("next", "下一批先补齐首批清单剩余七项", "**优先处理 next_download_7.csv 的七个目标。** 包括 M0 同比、人民币存贷款余额同比、新增人民币贷款累计、新增人民币存款累计、社融增量累计、社融存量。前三种统计口径中的水平、同比、当月、累计必须分别选择；社融中的人民币贷款分项也不能代替金融统计的新增人民币贷款总量。\n\n"
     "重导时选择原始单位和原始频率，空值留空，关闭换汇、插值、季度转月度与前向填充；保留代码、名称、单位、来源及导出时间。季度的房地产、个人购房、普惠贷款分项是额外变量，保留季度即可，优先级低于上述七项。\n\n"
     "reexport_10.csv 列出本批十个变换字段的原始代码。若方便，再重导 M0001383 和 M5525763 的不填零版本，以分清真实零与稀疏历史；无需逐月手填 PIT 标签。"),
    ("limits", "仍待确认的内容与使用边界", "**缺失原因和历史版本仍需另补证据。** 工作簿未保留可靠的导出时间字段，也没有逐期发布时间、首发版本或新旧口径说明。原始金融数据必须重新导出；文件里的计算记录只有操作参数，没有还原原值所需的全部汇率及原始空值标记。\n\n"
     "PIT_D 的真实可用时间从本项目首次归档时起算，未实现估计发布日期研究视图。全量验收仍为 FAIL：必需官方来源 4/5，可验证 A/B 占中国记录 58.3%。核心指标 16/16、历史深度 26/25 是现有验收程序包含 Wind D 的数值覆盖统计，不能理解为严格 PIT 完成；严格宽表仍为 259 行、47 个指标。原件可追溯率 100%，重复记录 0。"),
]
body = "# Wind 首批接入结果\n\n"
for ident, title, text in sections:
    body += f"## {title}\n\n{text}\n\n"
    if ident == "coverage":
        body += markdown_table(coverage_rows) + "\n\n"
body += ("## 复核文件\n\n- 当前值宽表：`wind_current_history_values.csv`（行索引为数据期，不是 as_of）。\n"
         "- 逐字段起止与变换：`field_profile.csv`；全表/排除单元格：`all_source_cells.csv`、`quarantined_cells.csv`。\n"
         "- 原件及时间：`manifest.json`；内嵌设置：`embedded_export_settings.json`；入库：`ingestion_result.json`。\n"
         "- 可执行复核：`review.ipynb`；`python scripts/review_wind_batch.py` 仅重新生成检查文件，在内存库验证，不写主库。\n"
         "- `--ingest` 才写主库；主库为追加写入，首次结果保留，重复运行另记 replay_result.json。\n")
body += ("\nHTML 打包状态：本次已尝试标准交付命令，但本机没有 Node/npm，未能生成 HTML。"
         "报告正文保存在本说明中，结构化输入保存在 `artifact.json`；检查与入库不依赖该渲染环境。\n")
(OUT / "README.md").write_text(body, encoding="utf-8")

source = {"id": "intake", "label": "Wind 首批工作簿及离线复核", "path": "reports/v2/wind_batch1/manifest.json",
          "query": {"engine": "Python", "language": "python", "description": "读取 SHA 归档的 数据.xlsx；检查每个指标和单元格，输出字段覆盖、内嵌操作配置、入库与全量验收。",
                    "tables_used": ["数据.xlsx / 中国_M1_同比", "field_profile.csv", "summary.json", "acceptance.txt"],
                    "executed_at": summary["reviewed_at"]}}
blocks = [{"id": "title", "type": "markdown", "body": "# Wind Batch Intake"}]
for ident, title, text in sections:
    blocks.append({"id": ident, "type": "markdown", "body": f"## {title}\n\n{text}", "sourceId": "intake"})
    if ident == "coverage":
        blocks.append({"id": "coverage_table_block", "type": "table", "tableId": "coverage_table"})
artifact = {"surface": "report", "manifest": {
    "version": 1, "surface": "report", "title": "Wind Batch Intake", "generatedAt": summary["reviewed_at"],
    "blocks": blocks, "sources": [source], "cards": [], "charts": [], "tables": [{
        "id": "coverage_table", "title": "本批接入数据期", "dataset": "coverage", "sourceId": "intake",
        "defaultSort": {"field": "Wind代码", "direction": "asc"},
        "columns": [{"field": k, "label": k, "type": "number" if k == "入库数据期" else "text"} for k in coverage_rows[0]],
    }]}, "snapshot": {"version": 1, "generatedAt": summary["reviewed_at"], "status": "partial",
                      "datasets": {"coverage": coverage_rows}, "accessIssues": []}, "sources": [source]}
(OUT / "artifact.json").write_text(json.dumps(artifact, ensure_ascii=False, indent=2), encoding="utf-8")
(OUT / "report_notes.json").write_text(json.dumps({
    "audience": "product stakeholders", "delivery": "html", "structure": "title, summary, findings+coverage, next steps, open questions+caveats",
    "visual_contract": "Exact lookup table: eight series, first/last accepted period and count; first period is not continuity. No color encoding. A chart would obscure individual dates and codes.",
    "sources": ["manifest.json", "field_profile.csv", "embedded_export_settings.json", "summary.json", "acceptance.txt"],
}, ensure_ascii=False, indent=2), encoding="utf-8")
print("Wrote README.md, next_download_7.csv and canonical artifact.json")

code_cells = [
    '''from pathlib import Path
import os, sys, json, hashlib, math
root = Path.cwd()
if not (root / 'src/macro_pit').exists():
    root = next(p for p in root.parents if (p / 'src/macro_pit').exists())
os.chdir(root)
sys.path.insert(0, str(root / 'src'))
import pandas as pd
from openpyxl import load_workbook
out = Path('reports/v2/wind_batch1')
manifest = json.loads((out / 'manifest.json').read_text(encoding='utf-8'))
assert hashlib.sha256(Path(manifest['path']).read_bytes()).hexdigest() == manifest['sha256']
summary = json.loads((out / 'summary.json').read_text(encoding='utf-8'))
assert summary['accepted_cells'] == 1958
print({k: summary[k] for k in ['input_rows', 'input_indicators', 'zero_cells', 'accepted_indicators', 'accepted_cells']})
''',
    '''# Independent comparison: read each accepted cell directly from the workbook.
# Do not use the intake parser or its mapping table for expected values.
wb = load_workbook(manifest['path'], data_only=True)
source_cells = pd.read_csv(out / 'all_source_cells.csv')
accepted_cells = source_cells[source_cells.status.eq('accepted')]
normalized = pd.read_csv(out / 'accepted_values_long.csv')
assert len(normalized) == len(accepted_cells) == 1958
assert not normalized.duplicated(['canonical_series_id', 'period']).any()
values = normalized.set_index(['source_series_id', 'period'])
for row in accepted_cells.itertuples():
    sheet = wb[row.sheet]
    cell = sheet[row.cell]
    source_unit = sheet.cell(4, cell.column).value
    wind_code = sheet.cell(5, cell.column).value
    period = sheet.cell(cell.row, 1).value.strftime('%Y-%m')
    assert wind_code == row.wind_code and period == row.period
    scale = {'%': 1, '\u4ebf\u7f8e\u5143': 0.1}[source_unit]
    assert math.isclose(values.loc[(wind_code, period), 'value'], cell.value * scale, rel_tol=0, abs_tol=1e-10)
wb.close()
assert normalized.pit_grade.eq('D').all() and normalized.source.eq('WIND').all()
assert normalized.release_at.isna().all()
print('1958/1958 source-cell, period, unit and value comparisons passed')
print(pd.read_csv(out / 'field_profile.csv')[['wind_code', 'accepted_first_period', 'accepted_last_period', 'accepted_cells']].to_string(index=False))
''',
    '''from datetime import datetime, timedelta
from macro_pit.db import get_connection, insert_observations
from macro_pit.pit import get_snapshot
from macro_pit.wind import review_workbook
records = review_workbook(manifest)['observations']
cutoff = datetime.fromisoformat(manifest['first_seen_at'])
with get_connection(':memory:') as conn:
    first = insert_observations(conn, records)
    replay = insert_observations(conn, records)
    assert first.inserted == 1958 and replay.inserted == 0 and replay.unchanged == 1958
    for mode in ['strict', 'loose']:
        assert get_snapshot(conn, cutoff.isoformat(), 'CN', mode).is_empty()
    assert get_snapshot(conn, (cutoff-timedelta(seconds=1)).isoformat(), 'CN', 'observed').is_empty()
    assert get_snapshot(conn, cutoff.isoformat(), 'CN', 'observed').height == 1958
print('Idempotency, strict/loose exclusion and first_seen boundary passed')
''',
    '''# Read-only confirmation of this frozen batch in the main database.
with get_connection('macro_pit_v2.duckdb', read_only=True) as conn:
    count, wrong = conn.execute("SELECT count(*), count(*) FILTER (WHERE pit_grade != 'D' OR release_at IS NOT NULL OR available_at != first_seen_at) FROM observation_vintage WHERE source='WIND' AND raw_sha256=?", [manifest['sha256']]).fetchone()
assert count == 1958 and wrong == 0
ingestion = json.loads((out / 'ingestion_result.json').read_text(encoding='utf-8'))
assert ingestion['strict_snapshot_unchanged'] and ingestion['strict_exports_unchanged']
identity = pd.read_csv(out / 'trade_identity_check.csv')
assert len(identity) == 260
assert identity.balance_residual_100mn_usd.abs().max() <= 0.02
print('Main DB batch count and grade checks passed; 260 trade identities within 0.02 hundred-million USD')
print('Strict export invariance is the saved ingestion-time check; future official backfills may legitimately change those files.')
''',
]
notebook = {
    "nbformat": 4, "nbformat_minor": 5,
    "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}},
    "cells": [{"cell_type": "markdown", "id": "intro", "metadata": {}, "source": [
        "# Wind 首批离线复核\n", "归档 SHA、1,958 个原始单元格独立核对、内存库重复/时间边界、主库只读确认。不会执行工作簿公式或写主库。\n"]}] + [
            {"cell_type": "code", "id": f"check-{i}", "metadata": {}, "source": code.splitlines(True),
             "execution_count": None, "outputs": []} for i, code in enumerate(code_cells, 1)],
}
(OUT / "review.ipynb").write_text(json.dumps(notebook, ensure_ascii=False, indent=2), encoding="utf-8")
